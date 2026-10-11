"""Transforma o relatório (lista de blocos) em DOCX ou PDF."""

import html
import io

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from revalli.modules.reports.document import Figure, Heading, PageBreak, Para, Report, Table

GRAY = RGBColor(0x57, 0x53, 0x4E)


# --- DOCX --------------------------------------------------------------------------------


def _shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    props.append(shd)


def _repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    props.append(el)


def to_docx(report: Report) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.2))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    width = section.page_width - section.left_margin - section.right_margin

    doc.add_heading(report.title, level=0)
    sub = doc.add_paragraph(report.subtitle)
    sub.runs[0].font.color.rgb = GRAY

    def caption(text: str) -> None:
        par = doc.add_paragraph()
        run = par.add_run(text)
        run.italic = True
        run.font.size = Pt(9)

    for b in report.blocks:
        if isinstance(b, Heading):
            doc.add_heading(b.text, level=b.level)
        elif isinstance(b, Para):
            par = doc.add_paragraph()
            run = par.add_run(b.text)
            if b.style == "note":
                run.italic = True
                run.font.color.rgb = GRAY
                run.font.size = Pt(9.5)
            elif b.style == "mono":
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif b.style == "label":
                run.bold = True
        elif isinstance(b, Table):
            caption(b.caption)
            table = doc.add_table(rows=1, cols=len(b.head))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, b.head, strict=True):
                cell.text = ""
                run = cell.paragraphs[0].add_run(text)
                run.bold = True
                run.font.size = Pt(9)
                _shade(cell, "EEEEEE")
            _repeat_header(table.rows[0])
            for r in b.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, r, strict=True):
                    cell.text = ""
                    cell.paragraphs[0].add_run(text).font.size = Pt(9)
            doc.add_paragraph()
        elif isinstance(b, Figure):
            doc.add_picture(io.BytesIO(b.png), width=_fit(b.png, width))
            caption(b.caption)
        elif isinstance(b, PageBreak):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _fit(png: bytes, width: int) -> int:
    """Figuras pequenas não são esticadas; as grandes ocupam a largura da página."""
    pix = pymupdf.Pixmap(png)
    # As figuras chegam em escala 2 (telas de alta densidade): 1 px = 1/192 polegada
    natural = int(pix.width / 192 * 914400)
    return min(width, natural)


# --- PDF ---------------------------------------------------------------------------------

CSS = """
body { font-family: sans-serif; font-size: 10pt; line-height: 1.35; }
h1 { font-size: 16pt; margin: 14pt 0 6pt 0; }
h2 { font-size: 13pt; margin: 12pt 0 4pt 0; }
h3 { font-size: 11pt; margin: 10pt 0 3pt 0; }
.title { font-size: 20pt; font-weight: bold; margin: 0 0 4pt 0; }
.subtitle { color: #57534e; margin: 0 0 10pt 0; }
p { margin: 0 0 5pt 0; }
.note { color: #57534e; font-style: italic; font-size: 9pt; }
.mono { font-family: monospace; font-size: 8.5pt; }
.label { font-weight: bold; }
.caption { font-style: italic; font-size: 8.5pt; margin: 6pt 0 3pt 0; }
table { border-collapse: collapse; width: 100%; margin: 0 0 8pt 0; }
th, td { border: 0.5pt solid #999; padding: 2pt 3pt; text-align: left; vertical-align: top;
         font-size: 8.5pt; }
th { border-bottom: 1pt solid #444; }
"""


def _segments(report: Report) -> list[str | Figure | PageBreak]:
    """HTML corrido, interrompido pelas figuras e quebras de página (posicionadas à parte)."""
    e = html.escape
    out: list[str | Figure | PageBreak] = []
    parts = [f'<p class="title">{e(report.title)}</p><p class="subtitle">{e(report.subtitle)}</p>']

    def flush() -> None:
        if parts:
            out.append("".join(parts))
            parts.clear()

    for b in report.blocks:
        if isinstance(b, Heading):
            level = min(3, b.level)
            parts.append(f"<h{level}>{e(b.text)}</h{level}>")
        elif isinstance(b, Para):
            text = e(b.text).replace("\n", "<br/>")
            parts.append(f'<p class="{b.style}">{text}</p>')
        elif isinstance(b, Table):
            head = "".join(f"<th>{e(h)}</th>" for h in b.head)
            rows = "".join(
                "<tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in b.rows
            )
            parts.append(
                f'<p class="caption">{e(b.caption)}</p><table><tr>{head}</tr>{rows}</table>'
            )
        else:
            flush()
            out.append(b)
    flush()
    return out


class _Pages:
    """Escreve blocos de HTML em sequência, abrindo páginas novas quando preciso."""

    def __init__(self) -> None:
        self.page = pymupdf.paper_rect("a4")
        self.body = self.page + (56, 56, -56, -56)
        self.buf = io.BytesIO()
        self.writer = pymupdf.DocumentWriter(self.buf)
        self.dev = None
        self.y = self.body.y0

    def new_page(self) -> None:
        if self.dev is not None:
            self.writer.end_page()
        self.dev = self.writer.begin_page(self.page)
        self.y = self.body.y0

    def place(self, story: pymupdf.Story, max_height: float | None = None) -> None:
        if self.dev is None:
            self.new_page()
        while True:
            bottom = self.body.y1 if max_height is None else min(self.body.y1, self.y + max_height)
            rect = pymupdf.Rect(self.body.x0, self.y, self.body.x1, bottom)
            more, filled = story.place(rect)
            story.draw(self.dev)
            if not more:
                self.y = pymupdf.Rect(filled).y1 + 4
                return
            self.new_page()

    def close(self) -> bytes:
        if self.dev is None:
            self.new_page()
        self.writer.end_page()
        self.writer.close()
        return self.buf.getvalue()


def to_pdf(report: Report) -> bytes:
    pages = _Pages()
    width = pages.body.width
    for seg in _segments(report):
        if isinstance(seg, str):
            pages.place(pymupdf.Story(html=seg, user_css=CSS))
        elif isinstance(seg, PageBreak):
            pages.new_page()
        else:
            pix = pymupdf.Pixmap(seg.png)
            # Figuras em escala 2: largura natural em pontos, limitada à largura útil
            w = min(width, pix.width * 72 / 192)
            h = w * pix.height / pix.width + 30  # figura + legenda
            if pages.dev is None or pages.y + h > pages.body.y1:
                pages.new_page()
            archive = pymupdf.Archive()
            archive.add(seg.png, "fig.png")
            story = pymupdf.Story(
                html=f'<img src="fig.png" width="{w:.0f}"/>'
                f'<p class="caption">{html.escape(seg.caption)}</p>',
                user_css=CSS,
                archive=archive,
            )
            pages.place(story)
    # Numera as páginas e grava os metadados
    doc = pymupdf.open("pdf", pages.close())
    for i, pg in enumerate(doc, start=1):
        pg.insert_text(
            (pg.rect.width - 70, pg.rect.height - 28), f"{i}/{doc.page_count}", fontsize=8
        )
    doc.set_metadata({"title": report.title, "creator": "ReValli", "producer": "ReValli"})
    return doc.tobytes(garbage=3, deflate=True)
